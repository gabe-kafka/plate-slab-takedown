"""Put the engine's tributary takedown into the firm's column load takedown workbook.

The firm's workbook (the 1025 Atlantic lineage: MASTER TRIB / MASTER FASCADE /
MASTER_KLL grids, C-BASE floor list, one C-(n) sheet per column that pulls its
row n+1 letter from the masters, and a Column Schedule) does the load combos,
reduction and column checks itself. This script only feeds it:

  MASTER TRIB, MASTER FASCADE, MASTER_KLL   column numbers across, slabs down
  C-BASE                                   elevations, floor labels, f'c, slab thickness
  C-(n)                                    one sheet per column (cloned from C-(1)),
                                           per-level SDL / facade psf / live load
  Column Schedule                          one pair of columns per column, on one sheet
  NOTES                                    column map and the assumptions used

usage: python scripts/fill_firm_takedown.py <engine column_load_takedown.xlsx>
           <firm template .xlsm/.xlsx> <levels.json> <out.xlsm>

levels.json: see tasks/1300_manhattan/takedown_levels.json. Levels are listed
top-down and land on C-BASE rows 93, 94, ...; the slab at a level is taken from
the engine floor named by slab_from. Each column segment carries the slab one
level up, as the firm's sheet is wired (C-(n) row r reads master row r-92).
"""
import json
import math
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

BASE_ROW = 93          # C-BASE / C-(n) row of the top level
MASTER_FIRST_ROW = 2   # master row of the top slab (= BASE_ROW + 1 - 92)
ROW_TO_MASTER = BASE_ROW + 1 - MASTER_FIRST_ROW  # C-(n) row r reads master row r - 92
LAST_ROW = 113
MASTER_REF = re.compile(
    r"((?:'MASTER TRIB'|'MASTER FASCADE'|MASTER_KLL)!)(\$?)([A-Z]{1,3})(\$?\d+)(?::(\$?)([A-Z]{1,3})(\$?\d+))?"
)

YELLOW = PatternFill("solid", fgColor="FFFF00")
F_NOTE = Font(name="Arial", size=9, italic=True, color="555555")
F_BOLD = Font(name="Arial", size=10, bold=True)
F_NORM = Font(name="Arial", size=10)


def label_key(lab):
    m = re.fullmatch(r"C(\d+)", str(lab))
    return (0, int(m.group(1))) if m else (1, str(lab))


def read_grid(ws):
    """Engine master sheet -> {floor: {label: value}}, labels in header order."""
    labels = [c.value for c in ws[1]][1:]
    grid = {}
    for r in range(2, ws.max_row + 1):
        floor = ws.cell(r, 1).value
        if floor is None:
            continue
        grid[str(floor)] = {lab: ws.cell(r, i + 2).value for i, lab in enumerate(labels) if lab is not None}
    return labels, grid


def put(ws, r, c, v):
    """ws.cell(r, c, None) leaves the cell alone; this always assigns."""
    ws.cell(r, c).value = v


def set_formula(cell, text):
    v = cell.value
    if isinstance(v, ArrayFormula):
        cell.value = ArrayFormula(v.ref, text)
    else:
        cell.value = text


def retarget_masters(ws, letter):
    """Point every master reference in ws at column `letter` of the masters."""
    def sub(m):
        rng = ""
        if m.group(6):
            rng = f":{m.group(5)}{letter}{m.group(7)}"
        return f"{m.group(1)}{m.group(2)}{letter}{m.group(4)}{rng}"

    for row in ws.iter_rows():
        for c in row:
            v = c.value
            text = v.text if isinstance(v, ArrayFormula) else v
            if isinstance(text, str) and "MASTER" in text:
                new = MASTER_REF.sub(sub, text)
                if new != text:
                    set_formula(c, new)


def clone_sheet(wb, src, title, before):
    """copy_worksheet plus the pieces it leaves behind (conditional formats,
    view, print setup), inserted before sheet `before`."""
    dst = wb.copy_worksheet(src)
    dst.title = title
    dst.conditional_formatting = ConditionalFormattingList()
    for rng in src.conditional_formatting:
        for rule in rng.rules:
            dst.conditional_formatting.add(str(rng.sqref), copy(rule))
    dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.freeze_panes = src.freeze_panes
    if src.print_area:
        dst.print_area = src.print_area.split("!")[-1]
    if src.print_title_rows:
        dst.print_title_rows = src.print_title_rows
    dst.sheet_properties.pageSetUpPr = copy(src.sheet_properties.pageSetUpPr)
    wb._sheets.remove(dst)
    wb._sheets.insert(wb._sheets.index(before), dst)
    return dst


SCHED_FIRST_BAND = 5     # schedule rows 5-8 show C-(n) row 93, 9-12 row 94, ...
SCHED_BAND_ROWS = 4
SCHED_BANDS = 21
SCHED_PATTERN_BAND = 17  # a band whose six cells are all intact in the template
SCHED_FND_ROW = 89
CELL_REF = re.compile(r"(\$?)([A-Z]{1,2})(\$?)(\d+)")


def rebuild_schedule_bands(ws, pair_cols, n_levels):
    """Rewrite the six formula cells of every band x column pair from one intact
    band. The firm's template has bands with formulas missing or typed over
    ("12x24" literals), which never showed on buildings that start lower.
    Unused bands are hidden and FND LOADS reads the lowest level's row."""
    r1 = SCHED_PATTERN_BAND
    p1, q1 = pair_cols[0], pair_cols[0] + 1
    pattern = []
    for (r, c) in ((r1, p1), (r1, q1), (r1 + 1, p1), (r1 + 1, q1), (r1 + 2, q1), (r1 + 3, q1)):
        v = ws.cell(r, c).value
        text = v.text if isinstance(v, ArrayFormula) else v
        if not (isinstance(text, str) and text.startswith("=")):
            raise SystemExit(f"Column Schedule {L(c)}{r} is not a formula; cannot use band {r1} as the pattern")
        pattern.append((r - r1, c - p1, text))
    src_letter = L(p1)

    def shifted(text, r0, p):
        def sub(m):
            col_abs, col, row_abs, row = m.group(1), m.group(2), m.group(3), int(m.group(4))
            if col == src_letter and not col_abs:
                return f"{L(p)}{row_abs}{row if row_abs else row - r1 + r0}"
            if col_abs and r1 <= row < r1 + SCHED_BAND_ROWS:
                return f"${col}{row_abs}{row - r1 + r0}"
            return m.group(0)
        return CELL_REF.sub(sub, text)

    for b in range(SCHED_BANDS):
        r0 = SCHED_FIRST_BAND + b * SCHED_BAND_ROWS
        for p in pair_cols:
            for dr, dc, text in pattern:
                cell = ws.cell(r0 + dr, p + dc)
                cell.value = ArrayFormula(cell.coordinate, shifted(text, r0, p))
        hidden = b >= n_levels
        for r in range(r0, r0 + SCHED_BAND_ROWS):
            ws.row_dimensions[r].hidden = hidden
    put(ws, SCHED_FND_ROW, 1, BASE_ROW + n_levels - 1)
    # FND LOADS shows only where the last band has a size: point it at the lowest level's band
    last_size_row = SCHED_FIRST_BAND + (n_levels - 1) * SCHED_BAND_ROWS + 1
    tmpl_size_row = SCHED_FIRST_BAND + (SCHED_BANDS - 1) * SCHED_BAND_ROWS + 1
    for p in pair_cols:
        cell = ws.cell(SCHED_FND_ROW, p)
        v = cell.value
        text = v.text if isinstance(v, ArrayFormula) else v
        if isinstance(text, str) and text.startswith("="):
            set_formula(cell, re.sub(rf"(?<![A-Z$]){L(p)}{tmpl_size_row}(?!\d)", f"{L(p)}{last_size_row}", text))


SCHED_FIRST_PAIR = 16    # column P: the template's first column pair (P = sizes, Q = loads)
SCHED_LAST_ROW = 90
SCHED_LABEL_COLS = "N:O"


def reletter(text, old, new):
    """Move relative references from column `old` to column `new` (P$1 -> BO$1, P86 -> BO86)."""
    def sub(m):
        if m.group(2) == old and not m.group(1):
            return f"{new}{m.group(3)}{m.group(4)}"
        return m.group(0)
    return CELL_REF.sub(sub, text)


def widen_schedule(ws, n_pairs):
    """Grow the template's 25 column pairs to n_pairs on the one sheet: pair k is a
    copy of pair P:Q (values relettered, styles, widths, merges), the right-hand
    label column moves to the end, print setup spans the width. Returns the
    size column of every pair."""
    tmpl_pairs = [c for c in range(SCHED_FIRST_PAIR, ws.max_column + 1, 2)
                  if str(ws.cell(4, c).value or "").isdigit()]
    right = tmpl_pairs[-1] + 2                      # the template's right-hand label column
    rows = range(1, SCHED_LAST_ROW + 1)
    right_cells = [(r, ws.cell(r, right).value, copy(ws.cell(r, right)._style)) for r in rows]
    right_width = ws.column_dimensions[L(right)].width
    right_merges = [(m.min_row, m.max_row) for m in ws.merged_cells.ranges if m.min_col == right]
    title = [m for m in ws.merged_cells.ranges if m.min_row == 3 and m.min_col < SCHED_FIRST_PAIR and m.max_col >= right]
    for m in list(ws.merged_cells.ranges):
        if m.min_col >= right or m in title:
            ws.unmerge_cells(str(m))
    pair_merges = [(m.min_row, m.max_row, m.min_col - SCHED_FIRST_PAIR, m.max_col - SCHED_FIRST_PAIR)
                   for m in ws.merged_cells.ranges if m.min_col in (SCHED_FIRST_PAIR, SCHED_FIRST_PAIR + 1)]
    src = [[(ws.cell(r, SCHED_FIRST_PAIR + d).value, copy(ws.cell(r, SCHED_FIRST_PAIR + d)._style)) for d in (0, 1)] for r in rows]
    widths = [ws.column_dimensions[L(SCHED_FIRST_PAIR + d)].width for d in (0, 1)]
    for c in range(right, ws.max_column + 1):
        for r in rows:
            put(ws, r, c, None)
    pairs = [SCHED_FIRST_PAIR + 2 * k for k in range(n_pairs)]
    for p in pairs[len(tmpl_pairs):]:
        for d in (0, 1):
            ws.column_dimensions[L(p + d)].width = widths[d]
            for r, cells in zip(rows, src):
                v, style = cells[d]
                text = v.text if isinstance(v, ArrayFormula) else v
                if isinstance(text, str) and text.startswith("="):
                    text = reletter(text, L(SCHED_FIRST_PAIR), L(p))
                    v = ArrayFormula(f"{L(p + d)}{r}", text) if isinstance(v, ArrayFormula) else text
                cell = ws.cell(r, p + d)
                cell.value = v
                cell._style = copy(style)
        for r0, r1, c0, c1 in pair_merges:
            ws.merge_cells(start_row=r0, end_row=r1, start_column=p + c0, end_column=p + c1)
    last = pairs[-1] + 2
    ws.column_dimensions[L(last)].width = right_width
    for r, v, style in right_cells:
        cell = ws.cell(r, last)
        cell.value = v
        cell._style = style
    for r0, r1 in right_merges:
        ws.merge_cells(start_row=r0, end_row=r1, start_column=last, end_column=last)
    if title:
        ws.merge_cells(start_row=3, end_row=3, start_column=title[0].min_col, end_column=last)
    ws.print_area = f"N3:{L(last)}{SCHED_LAST_ROW}"
    ws.print_title_cols = SCHED_LABEL_COLS
    ws.page_setup.fitToWidth = 0
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return pairs


def fill_master(ws, title_a1, numbers, slabs, values, assumed):
    """Clear the grid and write header, slab rows and values (None stays blank)."""
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 40), max_col=max(ws.max_column, len(numbers) + 1)):
        for c in row:
            c.value = None
            c.fill = PatternFill()
    ws["A1"] = title_a1
    for j, n in enumerate(numbers):
        ws.cell(1, j + 2, str(n))
    for i, slab in enumerate(slabs):
        r = MASTER_FIRST_ROW + i
        ws.cell(r, 1, slab)
        for j, n in enumerate(numbers):
            v = values[i][j]
            if v is not None:
                ws.cell(r, j + 2, v)
        if assumed[i]:
            for j in range(len(numbers) + 1):
                ws.cell(r, j + 1).fill = YELLOW
    ws.cell(MASTER_FIRST_ROW + len(slabs) + 1, 1,
            "Yellow rows: level not in the drawing set, areas of the engine floor named in NOTES stand in.").font = F_NOTE


def main(argv):
    if len(argv) != 4:
        print(__doc__)
        return 2
    engine_path, template_path, levels_path, out_path = map(Path, argv)
    cfg = json.loads(levels_path.read_text(encoding="utf-8"))
    levels = cfg["levels"]
    fc = cfg.get("fc_ksi", 6)
    fy = cfg.get("fy_ksi", 60)
    facade_psf = cfg.get("facade_psf", 35)
    kll_default = cfg.get("kll_default", 4)

    # ------------------------------------------------------------ engine data
    ewb = openpyxl.load_workbook(engine_path, data_only=True)
    trib_labels, trib = read_grid(ewb["MASTER TRIBUTARY AREA"])
    _, fasc = read_grid(ewb["FASCADE LENGTH"])
    _, kll = read_grid(ewb["MASTER KLL"])
    labels = sorted([lab for lab in trib_labels if lab is not None], key=label_key)
    numbers = list(range(1, len(labels) + 1))
    floors_of = {lab: [f for f in trib if trib[f].get(lab) is not None] for lab in labels}
    for f in {lv["slab_from"] for lv in levels if lv.get("slab_from")}:
        if f not in trib:
            raise SystemExit(f"levels.json names engine floor {f!r}; the engine has {sorted(trib)}")

    # Slabs carried: every level but the bottom one (the firm's sheet charges a
    # segment with the slab one level up; the bottom slab loads no column).
    slabs = levels[:-1]

    def slab_value(grid, lv, lab, default=None, as_int=True):
        f = lv.get("slab_from")
        if not f:
            return None
        v = grid.get(f, {}).get(lab)
        if v is None:
            return default if trib[f].get(lab) is not None else None
        return int(math.ceil(v)) if as_int else v

    trib_vals = [[slab_value(trib, lv, lab) for lab in labels] for lv in slabs]
    fasc_vals = [[slab_value(fasc, lv, lab, default=0) for lab in labels] for lv in slabs]
    kll_vals = [[slab_value(kll, lv, lab, default=kll_default) for lab in labels] for lv in slabs]
    assumed = [bool(lv.get("assumed")) for lv in slabs]
    slab_names = ["BULKHEAD" if lv["level"] == "BLKH" else lv["level"] for lv in slabs]

    # --------------------------------------------------------------- template
    wb = openpyxl.load_workbook(template_path, keep_vba=template_path.suffix.lower() == ".xlsm")
    fill_master(wb["MASTER TRIB"], "SLAB ", numbers, slab_names, trib_vals, assumed)
    fill_master(wb["MASTER FASCADE"], "Floor", numbers, slab_names, fasc_vals, assumed)
    fill_master(wb["MASTER_KLL"], "SLAB ", numbers, slab_names, kll_vals, assumed)

    # C-BASE: one row per level from BASE_ROW down; rows past the list go blank.
    cb = wb["C-BASE"]
    for i in range(LAST_ROW - BASE_ROW + 1):
        r = BASE_ROW + i
        if i < len(levels):
            lv = levels[i]
            put(cb, r, 2, lv["elev"])                 # ELEVATION
            put(cb, r, 5, lv["level"])                # FLOOR label
            put(cb, r, 10, fc)                        # f'c
            put(cb, r, 11, fy)                        # fy
            carried = levels[i - 1] if i > 0 else None  # slab this segment carries
            put(cb, r, 14, carried.get("slab_in", 8) if carried and carried.get("slab_from") else 0)
        else:
            put(cb, r, 2, None)
            put(cb, r, 5, None)

    # C-(n): per-level inputs on the base sheet, then clone for every column.
    base = wb["C-(1)"]
    for i in range(LAST_ROW - BASE_ROW + 1):
        r = BASE_ROW + i
        carried = levels[i - 1] if 0 < i < len(levels) else None
        if carried and carried.get("slab_from"):
            put(base, r, 17, carried.get("sdl", 20))         # Q  SUPERIMPOSED DEAD LOAD (PSF)
            put(base, r, 18, carried.get("facade_psf", facade_psf))  # R  FACADE LOAD (PSF)
            put(base, r, 19, 0)                               # S  TRANS LOAD DL
            put(base, r, 20, carried.get("ll", 40))          # T  LIVE LOAD (PSF)
            put(base, r, 21, 0)                               # U  TRANS LL
            put(base, r, 24, 0)                               # X  W (COMP)
            put(base, r, 25, 0)                               # Y  W (TEN)
        else:
            for col in (17, 18, 19, 20, 21, 24, 25):
                put(base, r, col, None)
    retarget_masters(base, L(2))

    existing = {ws.title: ws for ws in wb.worksheets if re.fullmatch(r"C-\(\d+\)", ws.title)}
    tail = wb["Tie Spacing Based 7.10.5"] if "Tie Spacing Based 7.10.5" in wb.sheetnames else wb.worksheets[-1]
    # drop template column sheets beyond the first: they are re-cloned from C-(1)
    for title, ws in existing.items():
        if title != "C-(1)":
            wb.remove(ws)
    for n in numbers[1:]:
        ws = clone_sheet(wb, base, f"C-({n})", tail)
        retarget_masters(ws, L(n + 1))

    # Column Schedule: one sheet, one column pair per column, floor labels read from the sheets.
    sched = wb["Column Schedule"]
    for r in range(5, 86, 4):  # floor label rows 5, 9, ..., 85
        if r > 9:              # 5 and 9 are the template's BULKHEAD / ROOF literals
            ref = f'INDIRECT("\'"&P$1&"\'!"&$B{r})'
            sched.cell(r, 14, f'=IF({ref}=0,"",{ref})')
    sched_cols = widen_schedule(sched, len(numbers))
    rebuild_schedule_bands(sched, sched_cols, len(levels))
    if sched["O5"].value:
        sched["O5"] = f"{int(fc * 1000)} PSI"
    for n, c in zip(numbers, sched_cols):
        put(sched, 4, c, str(n))

    # NOTES
    notes = wb.create_sheet("NOTES", 0)
    notes.sheet_view.showGridLines = False
    notes.column_dimensions["A"].width = 14
    notes.column_dimensions["B"].width = 22
    notes.column_dimensions["C"].width = 16
    for col in "DEFGH":
        notes.column_dimensions[col].width = 12
    notes.column_dimensions["I"].width = 90
    r = 1
    notes.cell(r, 1, f"{cfg.get('project', '')} - COLUMN LOAD TAKEDOWN").font = Font(name="Arial", size=12, bold=True)
    r += 1
    notes.cell(r, 1, f"Tributary areas, facade lengths and KLL from the tributary engine on {cfg.get('set', '')}; "
                     f"filled into the firm's takedown template by scripts/fill_firm_takedown.py. "
                     f"Yellow = assumption to confirm. Column sizes are the template's defaults; set them in each C-(n) sheet.").font = F_NOTE
    r += 2
    for j, h in enumerate(["Level", "Elev (ft)", "Slab from floor", "Assumed", "Slab (in)", "SDL (psf)", "LL (psf)", "Facade (psf)", "Note"]):
        notes.cell(r, j + 1, h).font = F_BOLD
    r += 1
    for lv in levels:
        vals = [lv["level"], lv.get("elev"), lv.get("slab_from"), "yes" if lv.get("assumed") else "", lv.get("slab_in"),
                lv.get("sdl"), lv.get("ll"), lv.get("facade_psf", facade_psf) if lv.get("slab_from") else None, lv.get("note", "")]
        for j, v in enumerate(vals):
            c = notes.cell(r, j + 1, v)
            c.font = F_NORM
            if lv.get("assumed"):
                c.fill = YELLOW
        r += 1
    r += 1
    notes.cell(r, 1, f"f'c {fc} ksi, fy {fy} ksi on every level (C-BASE). KLL default {kll_default} where the engine gives none. "
                     f"The bottom level's slab ({levels[-1]['level']}) is not carried by any column segment; add a level below it if there is a cellar.").font = F_NOTE
    r += 2
    for j, h in enumerate(["Column", "Engine label", "Floors with slab", "", "", "", "", "", "Note"]):
        notes.cell(r, j + 1, h).font = F_BOLD
    r += 1
    for n, lab in zip(numbers, labels):
        notes.cell(r, 1, n).font = F_NORM
        notes.cell(r, 2, lab).font = F_NORM
        fl = sorted(floors_of[lab], key=lambda f: -float(re.sub(r"\D", "", f) or 0))
        notes.cell(r, 3, ", ".join(fl)).font = F_NORM
        notes.cell(r, 3).alignment = Alignment(horizontal="left")
        if "UNLABELED" in str(lab):
            notes.cell(r, 9, "No column tag in the drawings; numbered after the tagged columns.").font = F_NOTE
        r += 1

    wb.active = wb.sheetnames.index("NOTES")
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "NOTES"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"{out_path}: {len(numbers)} columns, {len(levels)} levels, {len(wb.sheetnames)} sheets")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
